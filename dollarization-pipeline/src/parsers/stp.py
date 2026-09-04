"""Sao Tome and Principe: uses a single xlsx file,
'Depósitos Bancários por moeda e tipo_2018-2026.xlsx', distributed directly
from the Banco Central de S. Tomé e Príncipe (BCSTP) website's
'Banco Central > Estatísticas Monetárias e Financeiras' section.

This file is updated monthly, and sheet 'DEPÓSITOS BANCÁRIOS' holds a table
breaking down commercial-bank deposits by resident/non-resident x currency
(local currency/foreign currency) x sector (public/non-financial
corporations/households/nonprofits) (units: Milhões nDb, i.e. millions of the
new dobra following the 2018 redenomination [1000 STD=1 STN]). Since the file
itself starts at 2018-01, it naturally avoids the unit discontinuity caused by
the redenomination (data before that date is not in this file).

Row structure (labels in column C, values from column D onward monthly,
header dates on row 6):
    (1) Residentes                    <- total resident deposits (local
                                          currency+foreign currency) = used as TD
        (1.1) Moeda Nacional          <- residents, local currency
        (1.2) Moeda Estrangeira       <- residents, foreign currency = used as FCD
    (2) Não Residentes                <- non-resident deposits (excluded)
    TOTAL((1)+(2))                    <- residents+non-residents combined
                                          (excluded; this project targets
                                          residents only)

Label strings can vary slightly in leading whitespace/indentation from file to
file (e.g. ' (1) Residentes', '         (1.1) Moeda Nacional'), so we search
by matching only the number inside the parentheses via regex, rather than
relying on fixed row numbers. The header row is likewise found by searching
for "the first row where a lot of year-month-shaped datetime values appear"
(so it stays robust if the file layout shifts by a row or two in the future).

Note: (1) Residentes on its own does not represent "total deposits" in a
balance-of-payments sense — it refers only to resident deposits. If you want
the bank-wide deposit total including non-residents, use the TOTAL((1)+(2))
row instead; but since this project's FCD/TD indicators are meant to capture
"the share of resident foreign-currency deposits," TD is likewise kept on a
resident basis (row (1)).

Historical extension (per a 2026-08-19 user report): BCSTP separately
distributes 'Agregados Monetários_2001-2026.xlsx' (a monetary aggregates time
series, monthly from 2001-12). In sheet 'AGREGADOS MONETÁRIOS',
M3(=M2+foreign-currency deposits) - M0's 'Moeda em Circulação' (currency in
circulation) = TD, and the 'Depósitos em Moeda Estrangeira' row = FCD. Unlike
the deposit-detail file above, this file is on a **whole-banking-system
basis, not resident-only**, so the concept differs slightly (cross-check at
2018-01: ratio≈28.9% on this file vs. 26.57% on the detail file
[residents-only] — close in order of magnitude but not identical). So we only
use this monetary-aggregates file as a supplement to extend coverage before
2018-01, where the detail file has no data (the detail file's values win on
any overlapping month)."""

import re
from datetime import datetime, timezone
from io import BytesIO

import openpyxl
import pandas as pd

from src.utils.logger import get_logger

logger = get_logger(__name__)

FILE_URL = "__RENDER__"

_DETAIL_URL = (
    "https://www.bcstp.st/Upload/New_DOC/ES/"
    "Dep%C3%B3sitos%20Banc%C3%A1rios%20por%20moeda%20e%20tipo_2018-2026.xlsx"
)
_AGGREGATES_URL = (
    "https://www.bcstp.st/Upload/New_DOC/ES/"
    "Agregados%20Monet%C3%A1rios_2001-2026.xlsx"
)

_SHEET_NAME = "DEPÓSITOS BANCÁRIOS"
_TD_LABEL_RE = re.compile(r"^\(1\)\s*Residentes")
_FCD_LABEL_RE = re.compile(r"^\(1\.2\)\s*Moeda Estrangeira")

_AGG_SHEET_NAME = "AGREGADOS MONETÁRIOS"
_AGG_LABEL_COL = 2  # column B
_AGG_M3_LABEL_RE = re.compile(r"^M3\b")
_AGG_CURRENCY_LABEL_RE = re.compile(r"^Moeda em Circula")
_AGG_FCD_LABEL_RE = re.compile(r"^Dep[oó]sitos em Moeda Estrangeira")
_LABEL_COL = 3  # column C
_FIRST_DATA_COL = 4  # column D


def _find_header_row(ws) -> int:
    for r in range(1, ws.max_row + 1):
        date_count = sum(
            1
            for c in range(_FIRST_DATA_COL, ws.max_column + 1)
            if isinstance(ws.cell(row=r, column=c).value, datetime)
        )
        if date_count >= 3:
            return r
    raise ValueError("date header row not found")


def _find_label_row(ws, pattern: re.Pattern, max_row: int) -> int | None:
    for r in range(1, max_row + 1):
        label = ws.cell(row=r, column=_LABEL_COL).value
        if isinstance(label, str) and pattern.match(label.strip()):
            return r
    return None


def parse(content: bytes, country_code: str) -> pd.DataFrame:
    raise NotImplementedError("STP merges the two detail/aggregate xlsx files via render()")


def _parse_detail(content: bytes, country_code: str) -> pd.DataFrame:
    now = datetime.now(timezone.utc).isoformat()
    columns = ["country_code", "year", "period", "indicator", "value", "updated_at"]

    wb = openpyxl.load_workbook(BytesIO(content), data_only=True)
    if _SHEET_NAME not in wb.sheetnames:
        logger.warning("[%s] sheet '%s' not found (available sheets: %s)", country_code, _SHEET_NAME, wb.sheetnames)
        return pd.DataFrame(columns=columns)
    ws = wb[_SHEET_NAME]

    header_row = _find_header_row(ws)
    td_row = _find_label_row(ws, _TD_LABEL_RE, ws.max_row)
    fcd_row = _find_label_row(ws, _FCD_LABEL_RE, ws.max_row)

    if td_row is None or fcd_row is None:
        logger.warning("[%s] TD(%s)/FCD(%s) label row not found", country_code, td_row, fcd_row)
        return pd.DataFrame(columns=columns)

    # Only iterate up to the last column that actually has a value.
    last_col = _FIRST_DATA_COL - 1
    for c in range(_FIRST_DATA_COL, ws.max_column + 1):
        if ws.cell(row=header_row, column=c).value is not None:
            last_col = c

    rows = []
    for c in range(_FIRST_DATA_COL, last_col + 1):
        date_val = ws.cell(row=header_row, column=c).value
        if not isinstance(date_val, datetime):
            continue
        td_val = ws.cell(row=td_row, column=c).value
        fcd_val = ws.cell(row=fcd_row, column=c).value
        if td_val is None or fcd_val is None:
            continue

        year = date_val.year
        period = f"{year}-{date_val.month:02d}"
        td = round(float(td_val), 4)
        fcd = round(float(fcd_val), 4)
        ratio = round((fcd / td) * 100, 2) if td else None

        for indicator, value in (("FCD", fcd), ("TD", td), ("FCD_TD_RATIO", ratio)):
            if value is None:
                continue
            rows.append({
                "country_code": country_code,
                "year": year,
                "period": period,
                "indicator": indicator,
                "value": value,
                "updated_at": now,
            })

    return pd.DataFrame(rows, columns=columns)


def _find_agg_label_row(ws, pattern: re.Pattern) -> int | None:
    for r in range(1, ws.max_row + 1):
        label = ws.cell(row=r, column=_AGG_LABEL_COL).value
        if isinstance(label, str) and pattern.match(label.strip()):
            return r
    return None


def _parse_aggregates(content: bytes, country_code: str) -> pd.DataFrame:
    """'Agregados Monetários' monetary aggregates file — whole-banking-system
    basis (not resident-only).
    TD = M3 - Moeda em Circulação, FCD = Depósitos em Moeda Estrangeira."""
    now = datetime.now(timezone.utc).isoformat()
    columns = ["country_code", "year", "period", "indicator", "value", "updated_at"]

    wb = openpyxl.load_workbook(BytesIO(content), data_only=True)
    if _AGG_SHEET_NAME not in wb.sheetnames:
        logger.warning("[%s] sheet '%s' not found (available sheets: %s)", country_code, _AGG_SHEET_NAME, wb.sheetnames)
        return pd.DataFrame(columns=columns)
    ws = wb[_AGG_SHEET_NAME]

    header_row = None
    for r in range(1, min(10, ws.max_row) + 1):
        if sum(1 for c in range(3, ws.max_column + 1) if isinstance(ws.cell(row=r, column=c).value, datetime)) >= 3:
            header_row = r
            break
    m3_row = _find_agg_label_row(ws, _AGG_M3_LABEL_RE)
    currency_row = _find_agg_label_row(ws, _AGG_CURRENCY_LABEL_RE)
    fcd_row = _find_agg_label_row(ws, _AGG_FCD_LABEL_RE)
    if not (header_row and m3_row and currency_row and fcd_row):
        logger.warning(
            "[%s] rows not found in monetary aggregates file (header=%s m3=%s currency=%s fcd=%s)",
            country_code, header_row, m3_row, currency_row, fcd_row,
        )
        return pd.DataFrame(columns=columns)

    rows = []
    for c in range(3, ws.max_column + 1):
        date_val = ws.cell(row=header_row, column=c).value
        if not isinstance(date_val, datetime):
            continue
        m3 = ws.cell(row=m3_row, column=c).value
        currency = ws.cell(row=currency_row, column=c).value
        fcd_val = ws.cell(row=fcd_row, column=c).value
        if not all(isinstance(v, (int, float)) for v in (m3, currency, fcd_val)):
            continue

        year = date_val.year
        period = f"{year}-{date_val.month:02d}"
        td = round(m3 - currency, 4)
        fcd = round(float(fcd_val), 4)
        if td <= 0 or fcd < 0:
            continue
        ratio = round((fcd / td) * 100, 2)

        for indicator, value in (("FCD", fcd), ("TD", td), ("FCD_TD_RATIO", ratio)):
            rows.append({
                "country_code": country_code,
                "year": year,
                "period": period,
                "indicator": indicator,
                "value": value,
                "updated_at": now,
            })

    return pd.DataFrame(rows, columns=columns)


def render(target: dict) -> pd.DataFrame:
    import requests

    country_code = target["country_code"]
    headers = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/122.0.0.0 Safari/537.36"}

    frames = []
    try:
        resp = requests.get(_AGGREGATES_URL, headers=headers, timeout=90, verify=False)
        resp.raise_for_status()
        agg = _parse_aggregates(resp.content, country_code)
        if not agg.empty:
            logger.info("[%s] monetary aggregates %d rows (%s~%s)", country_code, len(agg), agg["period"].min(), agg["period"].max())
            frames.append(agg)
    except Exception as e:
        logger.warning("[%s] monetary aggregates file failed: %s", country_code, e)

    try:
        resp = requests.get(_DETAIL_URL, headers=headers, timeout=90, verify=False)
        resp.raise_for_status()
        detail = _parse_detail(resp.content, country_code)
        if not detail.empty:
            logger.info("[%s] detail %d rows (%s~%s)", country_code, len(detail), detail["period"].min(), detail["period"].max())
            frames.append(detail)
    except Exception as e:
        logger.warning("[%s] detail file failed: %s", country_code, e)

    if not frames:
        return pd.DataFrame(columns=["country_code", "year", "period", "indicator", "value", "updated_at"])

    # Concat monetary aggregates first, then detail file → with
    # drop_duplicates(keep='last'), the resident-basis detail value wins for
    # any overlapping month (2018-01~).
    out = (
        pd.concat(frames, ignore_index=True)
        .drop_duplicates(subset=["period", "indicator"], keep="last")
        .sort_values(["period", "indicator"])
        .reset_index(drop=True)
    )
    logger.info("[%s] merged %d rows (%s~%s)", country_code, len(out), out["period"].min(), out["period"].max())
    return out
