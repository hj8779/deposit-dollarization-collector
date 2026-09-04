"""Bangladesh: Bangladesh Bank's consolidated 'Time Series Data' XLSX (a direct link embedded
in the econdata page with no special notice: /econdata/time_series_data1972-2024.xlsx). A
direct GET on the download URL returns bot-protection challenge (TSPD) HTML, so Playwright
must be used to actually capture the download event (note that page.goto() throwing an
exception when the download starts is expected Playwright behavior and must be handled).

Inside sheet 'Table IA', five sub-tables from different eras are stacked vertically, and each
sub-table has a completely different column layout (more columns appear in later years). The
'foreign currency deposit' item only exists in the third through fifth sub-tables (the
first/second sub-tables, covering 1971-72~1987-88, don't have this item at all), and even the
label differs by era:
    Sub-table 3 (1988-89~2019-20): 'Foreign Currency Deposit Liabilities'
    Sub-table 4 (2020-21):          'Short Term FC Deposit Liabilities'
    Sub-table 5 (2021-22~2023-24):  'Short Term FC Deposit Liabilities'

Caution: all three sub-tables happen to place this item in the same column (column 39, Excel
column AM), but **the column index must not be hardcoded.** In the second sub-table (DMBs
Borrowings, 1972-73~1987-88), column 39 happens to be 'From Inter-Banks' (interbank borrowing,
unrelated to foreign currency deposits), so if you trust the index alone across the whole
range, the 1974-1987 data gets contaminated with completely wrong indicator values (this is
in fact what happened in an earlier implementation, and was caught while filtering out bad
values). So for every sub-table, the header text (a cell containing 'Foreign Currency
Deposit'/'FC Deposit') is located directly and its column number is used, and sub-tables
without that header are skipped.

Because the period follows Bangladesh's fiscal year (July to June of the following year,
formatted 'YYYY-YY' or, more recently, with a 'P' = provisional suffix), it cannot be mapped
to a precise month, so it's recorded as '{year}-Annual' based on the starting year.

TD (total deposits) = the 'Total Deposit Liabilities (37+38)' column. In all three sub-tables
this sits immediately to the right of the FCD header (col+1), and per the formula it's exactly
col37 (DMBs Deposits, local-currency deposits) + col38 (FCD).
"""

import re
from datetime import datetime, timezone
from io import BytesIO

import pandas as pd

from src.collectors.base import INDICATOR, INDICATOR_TD
from src.utils.logger import get_logger

logger = get_logger(__name__)

FILE_URL = "__RENDER__"

_SHEET_NAME = "Table IA"
_FY_RE = re.compile(r"^(\d{4})-\d{2}P?\s*$")
_HEADER_RE = re.compile(r"foreign currency deposit|fc deposit", re.I)


def parse(content: bytes, country_code: str) -> pd.DataFrame:
    raise NotImplementedError("BGD is handled via render() (Playwright required for the download)")


def _download_xlsx(country_code: str) -> bytes:
    from playwright.sync_api import sync_playwright

    file_url = "https://www.bb.org.bd/econdata/time_series_data1972-2024.xlsx"

    with sync_playwright() as p:
        browser = p.chromium.launch()
        context = browser.new_context(
            accept_downloads=True,
            user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                       "(KHTML, like Gecko) Chrome/120.0 Safari/537.36",
        )
        page = context.new_page()
        try:
            with page.expect_download(timeout=30000) as dl_info:
                try:
                    page.goto(file_url, timeout=30000)
                except Exception:
                    pass  # it's expected Playwright behavior for goto() to throw once the download starts
            download = dl_info.value
            path = f"/tmp/{country_code.lower()}_time_series.xlsx"
            download.save_as(path)
        finally:
            browser.close()

    return open(path, "rb").read()


def render(target: dict) -> pd.DataFrame:
    import openpyxl

    country_code = target["country_code"]
    now = datetime.now(timezone.utc).isoformat()

    content = _download_xlsx(country_code)
    wb = openpyxl.load_workbook(BytesIO(content), data_only=True)
    ws = wb[_SHEET_NAME]

    # 1) Find every (row, col) where the 'Foreign Currency Deposit' / 'FC Deposit' header appears.
    header_hits = []
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str) and _HEADER_RE.search(v):
                header_hits.append((r, c))

    if not header_hits:
        logger.warning("[%s] could not find the 'Foreign Currency Deposit' header", country_code)
        return pd.DataFrame(columns=["country_code", "year", "period", "indicator", "value", "updated_at"])

    # 1-1) Also find the 'Total Deposit Liabilities' header (the column right after the FCD header).
    td_col_by_header_row: dict[int, int] = {}
    for header_row, col in header_hits:
        next_header = ws.cell(row=header_row, column=col + 1).value
        if isinstance(next_header, str) and "total" in next_header.lower() and "deposit" in next_header.lower():
            td_col_by_header_row[header_row] = col + 1

    # 2) Walk down below each header, collecting values from that column as long as fiscal-year
    #    rows (col1) are found. Stop at footnote rows like 'Note:'/'Source:'.
    rows = []
    for header_row, col in header_hits:
        td_col = td_col_by_header_row.get(header_row)
        r = header_row + 1
        while r <= ws.max_row:
            label = ws.cell(row=r, column=1).value
            if isinstance(label, str) and label.strip().lower().startswith(("note", "source")):
                break
            if isinstance(label, str):
                m = _FY_RE.match(label.strip())
                if m:
                    year = int(m.group(1))
                    period = f"{year}-Annual"
                    value = ws.cell(row=r, column=col).value
                    if isinstance(value, (int, float)):
                        rows.append({
                            "country_code": country_code,
                            "year": year,
                            "period": period,
                            "indicator": INDICATOR,
                            "value": round(float(value), 2),
                            "updated_at": now,
                        })
                    if td_col is not None:
                        td_value = ws.cell(row=r, column=td_col).value
                        if isinstance(td_value, (int, float)):
                            rows.append({
                                "country_code": country_code,
                                "year": year,
                                "period": period,
                                "indicator": INDICATOR_TD,
                                "value": round(float(td_value), 2),
                                "updated_at": now,
                            })
            r += 1
            if r - header_row > 60:  # safety guard: no sub-table should be this long
                break

    if not rows:
        return pd.DataFrame(columns=["country_code", "year", "period", "indicator", "value", "updated_at"])

    df = pd.DataFrame(rows).drop_duplicates(subset=["period", "indicator"], keep="last")
    return df.sort_values("period").reset_index(drop=True)
