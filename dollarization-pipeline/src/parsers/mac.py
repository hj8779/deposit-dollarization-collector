"""Macao: sheet 'Resident Deposits', header row13, value=residentDepositsTotal - residentDepositsMOP.

TD (total deposits) = column3 (residentDepositsTotal) itself (all resident deposits, regardless of currency)."""

from datetime import datetime, timezone
from io import BytesIO

import openpyxl
import pandas as pd

from src.collectors.base import INDICATOR, INDICATOR_TD

FILE_URL = None  # uses target['source_url'] as-is


def parse(content: bytes, country_code: str) -> pd.DataFrame:
    wb = openpyxl.load_workbook(BytesIO(content), data_only=True)
    ws = wb["Resident Deposits"]
    now = datetime.now(timezone.utc).isoformat()
    rows = []
    current_year = None
    for r in range(14, ws.max_row + 1):
        year = ws.cell(row=r, column=1).value
        month = ws.cell(row=r, column=2).value
        total = ws.cell(row=r, column=3).value
        mop = ws.cell(row=r, column=4).value
        if year is not None:
            current_year = year
        if current_year is None or month is None:
            continue
        if not isinstance(total, (int, float)) or not isinstance(mop, (int, float)):
            continue
        period = f"{current_year}-{int(month):02d}"
        rows.append({
            "country_code": country_code,
            "year": current_year,
            "period": period,
            "indicator": INDICATOR,
            "value": float(total) - float(mop),
            "updated_at": now,
        })
        rows.append({
            "country_code": country_code,
            "year": current_year,
            "period": period,
            "indicator": INDICATOR_TD,
            "value": float(total),
            "updated_at": now,
        })
    return pd.DataFrame(rows)
