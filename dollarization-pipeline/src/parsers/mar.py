"""Morocco: Bank Al-Maghrib "Séries statistiques monétaires" — 'AGRÉGATS DE MONNAIE' xlsx.

Page: https://www.bkam.ma/Statistiques/Statistiques-monetaires/Series-statistiques-monetaires
The document ID in the download link (e.g.
/content/download/632818/7098245/6-AgregatsDeMonnaie.xlsx) can change with each
revision, so we scrape the page every time to find the current link.

Sheet 'Feuil1': row 3 is the month-end date header (starting 2001-12), and each
subsequent row is an indicator. Rows are looked up by label, not by row number:
  'Dépôts en devises' — footnote (3) "Dépôts à vue et à terme en devises auprès des banques"
      = foreign-currency demand + time deposits held at banks = FCD.
  'Monnaie scripturale' — sum of demand deposits held at BAM/banks/post office
      (CCP)/Treasury (the non-currency-in-circulation part of M1) = local-currency
      demand deposits.
  'Placements à vue' — savings-type instant-access deposits (quasi-money liquid holdings).
  'Comptes à terme et bons de caisse auprès des banques' — time deposits at banks.
  'Autres dépôts' — other deposits.

TD = Monnaie scripturale + Placements à vue + Comptes à terme... + Dépôts en devises
   + Autres dépôts. 'Titres OPCVM monétaires' (MMF shares), 'Valeurs données en
   pension' (repos), 'Certificats de dépôts' (NCDs), and 'Dépôts à terme auprès du
   Trésor' are excluded because they aren't pure deposits or are frequently marked
   not available (ND) — consistent with the NID/repo exclusion convention used in
   other country parsers.
   Cross-check: the 'Autres actifs Monétaires' row exactly matches the sum of the
   excluded items above plus Dépôts en devises + Comptes à terme + Autres dépôts
   (structure verified).

Monthly, 2001-12 to present. Unit: MDH (million dirhams)."""

from __future__ import annotations

import re
from datetime import datetime, timezone
from io import BytesIO

import openpyxl
import pandas as pd
import requests
import urllib3

from src.collectors.base import INDICATOR, INDICATOR_TD
from src.utils.logger import get_logger

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

logger = get_logger(__name__)

FILE_URL = "__RENDER__"

_LIST_PAGE = "https://www.bkam.ma/Statistiques/Statistiques-monetaires/Series-statistiques-monetaires"
_BASE = "https://www.bkam.ma"
_XLSX_LINK_RE = re.compile(r'href="([^"]*AgregatsDeMonnaie\.xlsx)"', re.I)

_HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/122.0.0.0 Safari/537.36"
    ),
}

_SCRIPTURAL_LABEL = "monnaie scripturale"
_SIGHT_SAVINGS_LABEL = "placements à vue"
_TERM_LABEL = "comptes à terme et bons de caisse"
_FCD_LABEL = "dépôts en devises"
_OTHER_LABEL = "autres dépôts"


def parse(content: bytes, country_code: str) -> pd.DataFrame:
    raise NotImplementedError("MAR fetches the xlsx via render()")


def _empty() -> pd.DataFrame:
    return pd.DataFrame(columns=["country_code", "year", "period", "indicator", "value", "updated_at"])


def _find_xlsx_url() -> str | None:
    resp = requests.get(_LIST_PAGE, headers=_HEADERS, timeout=30, verify=False)
    resp.raise_for_status()
    match = _XLSX_LINK_RE.search(resp.text)
    if not match:
        return None
    href = match.group(1)
    return href if href.startswith("http") else _BASE + href


def _find_row(ws, label: str) -> int | None:
    for r in range(1, ws.max_row + 1):
        cell = ws.cell(row=r, column=1).value
        if isinstance(cell, str) and cell.strip().lower().startswith(label):
            return r
    return None


def render(target: dict) -> pd.DataFrame:
    country_code = target["country_code"]
    xlsx_url = _find_xlsx_url()
    if not xlsx_url:
        logger.warning("[%s] could not find the AgregatsDeMonnaie.xlsx link on the listing page", country_code)
        return _empty()

    resp = requests.get(xlsx_url, headers=_HEADERS, timeout=60, verify=False)
    resp.raise_for_status()
    return _parse_workbook(resp.content, country_code)


def _parse_workbook(content: bytes, country_code: str) -> pd.DataFrame:
    wb = openpyxl.load_workbook(BytesIO(content), data_only=True)
    ws = wb["Feuil1"]

    scriptural_row = _find_row(ws, _SCRIPTURAL_LABEL)
    sight_savings_row = _find_row(ws, _SIGHT_SAVINGS_LABEL)
    term_row = _find_row(ws, _TERM_LABEL)
    fcd_row = _find_row(ws, _FCD_LABEL)
    other_row = _find_row(ws, _OTHER_LABEL)
    if not all([scriptural_row, sight_savings_row, term_row, fcd_row, other_row]):
        logger.error(
            "[%s] could not find required rows: scriptural=%s sight_savings=%s term=%s fcd=%s other=%s",
            country_code, scriptural_row, sight_savings_row, term_row, fcd_row, other_row,
        )
        return _empty()

    now = datetime.now(timezone.utc).isoformat()
    rows = []
    for c in range(2, ws.max_column + 1):
        date = ws.cell(row=3, column=c).value
        if not isinstance(date, datetime):
            continue

        fcd = ws.cell(row=fcd_row, column=c).value
        scriptural = ws.cell(row=scriptural_row, column=c).value
        sight_savings = ws.cell(row=sight_savings_row, column=c).value
        term = ws.cell(row=term_row, column=c).value
        other = ws.cell(row=other_row, column=c).value
        parts = (fcd, scriptural, sight_savings, term, other)
        if any(not isinstance(v, (int, float)) for v in parts):
            continue

        fcd = float(fcd)
        td = sum(float(v) for v in parts)
        if td <= 0:
            continue

        period = f"{date.year}-{date.month:02d}"
        ratio = round((fcd / td) * 100, 4)
        for indicator, value in ((INDICATOR, round(fcd, 4)), (INDICATOR_TD, round(td, 4)), ("FCD_TD_RATIO", ratio)):
            rows.append({
                "country_code": country_code, "year": date.year, "period": period,
                "indicator": indicator, "value": value, "updated_at": now,
            })

    if not rows:
        return _empty()

    out = (
        pd.DataFrame(rows)
        .drop_duplicates(subset=["period", "indicator"], keep="last")
        .sort_values(["period", "indicator"])
        .reset_index(drop=True)
    )
    logger.info("[%s] %d rows (%s~%s)", country_code, len(out), out["period"].min(), out["period"].max())
    return out
