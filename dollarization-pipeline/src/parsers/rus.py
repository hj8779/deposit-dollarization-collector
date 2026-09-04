"""Russia: CBR 'Depository corporations survey' (survey_dc_new_e.xlsx), sheet 'Short Form'.
row12='Deposits of households in foreign currency', row13='...nonfinancial corporations...',
row14='...other financial corporations...' — the sum of these three rows is total FCD.
The header (row1) uses monthly labels in 'Mon, YYYY' format.

TD (total deposits) = Transferable deposits (row4+5+6, local currency) + Other
deposits (row8+9+10, local currency) + FCD (row12+13+14). Verified against actual
figures that the identities M1(row7)=M0(row3)+Transferable, M2(row11)=M1+Other,
M2X(row2)=M2+FCD+Certificates(row15) hold (e.g. M1 869362=M0 418872+100410+337313+12767)."""

from datetime import datetime, timezone
from io import BytesIO

import openpyxl
import pandas as pd

from src.collectors.base import INDICATOR, INDICATOR_TD

FILE_URL = "https://www.cbr.ru/vfs/eng/statistics/credit_statistics/survey/survey_dc_new_e.xlsx"

_MONTH_MAP = {
    "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
    "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12,
}
_FCD_ROWS = [12, 13, 14]  # households / nonfinancial corporations / other financial corporations
_TD_ROWS = [4, 5, 6, 8, 9, 10, 12, 13, 14]  # transferable + other (local) + FCD


def parse(content: bytes, country_code: str) -> pd.DataFrame:
    wb = openpyxl.load_workbook(BytesIO(content), data_only=True)
    ws = wb["Short Form"]
    now = datetime.now(timezone.utc).isoformat()

    rows = []
    for c in range(2, ws.max_column + 1):
        label = ws.cell(row=1, column=c).value
        if not label:
            continue
        try:
            month_str, year_str = str(label).split(",")
            month = _MONTH_MAP[month_str.strip().lower()]
            year = int(year_str.strip())
        except (ValueError, KeyError):
            continue
        period = f"{year}-{month:02d}"

        fcd_values = [ws.cell(row=r, column=c).value for r in _FCD_ROWS]
        if all(isinstance(v, (int, float)) for v in fcd_values):
            rows.append({
                "country_code": country_code,
                "year": year,
                "period": period,
                "indicator": INDICATOR,
                "value": float(sum(fcd_values)),
                "updated_at": now,
            })

        td_values = [ws.cell(row=r, column=c).value for r in _TD_ROWS]
        if all(isinstance(v, (int, float)) for v in td_values):
            rows.append({
                "country_code": country_code,
                "year": year,
                "period": period,
                "indicator": INDICATOR_TD,
                "value": float(sum(td_values)),
                "updated_at": now,
            })
    return pd.DataFrame(rows)
