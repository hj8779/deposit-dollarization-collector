"""Botswana: Bank of Botswana 'Publications' 목록(/publications?page=0,1,2,...)에 매달
'Botswana Economic and Financial Statistics - {Month} {Year}' xlsx가 게시된다. 목록은
EFS 외에도 다른 온갖 간행물이 섞여 있고 페이지네이션 번호가 계속 늘어나므로, 페이지 0의
pager에서 마지막 페이지 번호를 읽어 그 범위를 전부 순회하며 EFS 항목만 골라낸다.

각 파일의 시트 '3.16'(COMMERCIAL BANKS: FOREIGN CURRENCY ACCOUNTS AND TOTAL DEPOSITS)에
통화별(US$/GBP/ZAR/EUR/Other) 외화계좌를 풀라 환산액으로 합친 'Total Pula equivalent' 열이
FCD, 'Deposits (Pula)' 열이 TD(전체 통화 합계)다. 컬럼은 6/7행 헤더 텍스트(2줄에 걸쳐
쪼개져 있음, 예: 6행 'Total Pula' + 7행 'equivalent')로 매번 동적으로 찾는다.

이 표는 파일마다 최근 ~10~13년 롤링 윈도우만 담고 있다(예: 2024-01 발행분은 2013년부터,
2026-06 발행분은 2016년부터). 그래서 발행월이 흩어진 여러 파일을 모아야 전체 기간을
커버할 수 있다 - 'publications' 목록(JS 페이지네이션)에는 2024-01 이전 EFS 파일이
보이지 않지만, 목록에 안 뜨는 구형 직접 링크 하나(BFS-JAN-2015.xls, 사용자 제공)가
2004~2014년치를 담고 있어 이것도 함께 받는다. 이 파일은 옛 바이너리 .xls(OLE) 포맷이라
openpyxl이 아닌 xlrd로 읽어야 하고, 헤더가 3줄(대분류/통화단위 두 줄이 최신 xlsx보다
한 줄 더 나뉨)에 걸쳐 있어 별도 파싱 함수(_parse_legacy_xls)를 쓴다.

연도별 행 구조: 앞부분(예: 2013~2021)은 한 해에 한 행(월 없이 연도만, End-of-year=12월
값으로 취급), 이후로는 분기별(Mar/Jun/Sep/Dec), 최근에는 매월(Jan~Dec)로 세분화된다.
행의 연도 칸(1열)이 비어 있으면 바로 위 값을 그대로 이어받는다(forward-fill).

여러 파일에 걸쳐 겹치는 기간은 그 파일 자체가 커버하는 마지막 기간(=발행월)이 더 최신인
파일 쪽 값을 채택한다(central bank가 이후 개정치를 반영했을 수 있어 최신판 우선).
"""

import re
from datetime import datetime, timezone

import pandas as pd
import requests

from src.collectors.base import INDICATOR
from src.utils.logger import get_logger

logger = get_logger(__name__)

FILE_URL = "__RENDER__"

LIST_URL = "https://www.bankofbotswana.bw/publications"
LEGACY_XLS_URL = "https://www.bankofbotswana.bw/sites/default/files/publications/BFS-JAN-2015.xls"
_HEADERS = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0 Safari/537.36"}

_SHEET_NAME = "3.16"
_GROUP_ROW = 6
_UNIT_ROW = 7
_DATA_START_ROW = 8
_YEAR_COL = 1
_MONTH_COL = 2
_FCD_HEADER_RE = re.compile(r"total\s*pula", re.I)
_TD_HEADER_RE = re.compile(r"deposits", re.I)
_TD_UNIT_RE = re.compile(r"\(pula\)", re.I)

_MONTHS = {
    "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
    "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12,
}


def parse(content: bytes, country_code: str) -> pd.DataFrame:
    raise NotImplementedError("BWA는 render()를 통해 처리한다 (게시글 목록을 순회해야 함)")


def _last_page_number() -> int:
    from playwright.sync_api import sync_playwright

    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page()
        page.goto(f"{LIST_URL}?page=0", timeout=60000)
        page.wait_for_timeout(1500)
        html = page.content()
        browser.close()

    pages = [int(m) for m in re.findall(r'href="\?page=(\d+)"', html)]
    return max(pages) if pages else 0


def _collect_efs_links() -> dict[str, str]:
    from playwright.sync_api import sync_playwright

    last_page = _last_page_number()
    logger.info("[BWA] publications 목록 마지막 페이지 번호: %d", last_page)

    links: dict[str, str] = {}
    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page()
        for pg in range(0, last_page + 1):
            page.goto(f"{LIST_URL}?page={pg}", timeout=60000)
            page.wait_for_timeout(500)
            items = page.eval_on_selector_all(
                "a",
                "els => els.map(e => ({href: e.href, text: e.innerText}))",
            )
            for item in items:
                href, text = item["href"], (item["text"] or "").strip()
                if "Economic and Financial Statistics" in text and href.lower().endswith(".xlsx"):
                    links[href] = text
        browser.close()

    return links


def _find_fcd_td_columns(ws) -> tuple[int | None, int | None]:
    fcd_col = td_col = None
    for c in range(1, ws.max_column + 1):
        group = ws.cell(row=_GROUP_ROW, column=c).value or ""
        unit = ws.cell(row=_UNIT_ROW, column=c).value or ""
        combined = f"{group} {unit}"
        if _FCD_HEADER_RE.search(combined) and "equivalent" in unit.lower():
            fcd_col = c
        elif _TD_HEADER_RE.search(group) and _TD_UNIT_RE.search(unit):
            td_col = c
    return fcd_col, td_col


def _parse_workbook(content: bytes, country_code: str) -> pd.DataFrame:
    import openpyxl
    from io import BytesIO

    now = datetime.now(timezone.utc).isoformat()
    wb = openpyxl.load_workbook(BytesIO(content), data_only=True)
    if _SHEET_NAME not in wb.sheetnames:
        return pd.DataFrame(columns=["country_code", "year", "period", "indicator", "value", "updated_at"])
    ws = wb[_SHEET_NAME]

    fcd_col, td_col = _find_fcd_td_columns(ws)
    if fcd_col is None:
        logger.warning("[%s] 'Total Pula equivalent' 컬럼을 찾지 못함", country_code)
        return pd.DataFrame(columns=["country_code", "year", "period", "indicator", "value", "updated_at"])

    rows = []
    current_year = None
    for r in range(_DATA_START_ROW, ws.max_row + 1):
        year_cell = ws.cell(row=r, column=_YEAR_COL).value
        month_cell = ws.cell(row=r, column=_MONTH_COL).value
        if isinstance(year_cell, (int, float)):
            current_year = int(year_cell)
        if current_year is None:
            continue

        fcd_val = ws.cell(row=r, column=fcd_col).value
        if not isinstance(fcd_val, (int, float)):
            continue

        if isinstance(month_cell, str) and month_cell.strip().lower()[:3] in _MONTHS:
            month = _MONTHS[month_cell.strip().lower()[:3]]
        elif isinstance(year_cell, (int, float)):
            month = 12  # 월 표기 없이 연도만 있는 행 = 그 해 12월(End of Period) 값
        else:
            continue

        period = f"{current_year}-{month:02d}"
        rows.append({
            "country_code": country_code, "year": current_year, "period": period,
            "indicator": INDICATOR, "value": round(float(fcd_val), 4), "updated_at": now,
        })

        td_val = ws.cell(row=r, column=td_col).value if td_col else None
        if isinstance(td_val, (int, float)):
            rows.append({
                "country_code": country_code, "year": current_year, "period": period,
                "indicator": "TD", "value": round(float(td_val), 4), "updated_at": now,
            })

    return pd.DataFrame(rows)


_LEGACY_YEAR_RE = re.compile(r"^(\d{4})")


def _parse_legacy_xls(content: bytes, country_code: str) -> pd.DataFrame:
    import xlrd
    from io import BytesIO

    now = datetime.now(timezone.utc).isoformat()
    wb = xlrd.open_workbook(file_contents=content)
    if _SHEET_NAME not in wb.sheet_names():
        return pd.DataFrame(columns=["country_code", "year", "period", "indicator", "value", "updated_at"])
    ws = wb.sheet_by_name(_SHEET_NAME)

    # 헤더가 3줄(대분류/통화단위가 최신 xlsx보다 한 줄 더 나뉨)에 걸쳐 있어, 각 열의
    # 헤더 후보 3줄을 합쳐서 텍스트로 찾는다(고정 인덱스 대신 헤더 텍스트 탐색 원칙 유지).
    header_rows = range(4, 7)  # 0-indexed
    data_start_row = 7

    fcd_col = td_col = None
    for c in range(ws.ncols):
        combined = " ".join(str(ws.cell_value(r, c)) for r in header_rows if ws.cell_value(r, c))
        if _FCD_HEADER_RE.search(combined) and "equivalent" in combined.lower():
            fcd_col = c
        elif _TD_HEADER_RE.search(combined) and _TD_UNIT_RE.search(combined):
            td_col = c

    if fcd_col is None:
        logger.warning("[%s] 레거시 xls에서 'Total Pula equivalent' 컬럼을 찾지 못함", country_code)
        return pd.DataFrame(columns=["country_code", "year", "period", "indicator", "value", "updated_at"])

    rows = []
    current_year = None
    for r in range(data_start_row, ws.nrows):
        year_cell = ws.cell_value(r, 0)
        month_cell = ws.cell_value(r, 1)

        m = _LEGACY_YEAR_RE.match(str(year_cell).strip()) if year_cell != "" else None
        if isinstance(year_cell, float):
            current_year = int(year_cell)
        elif m:
            current_year = int(m.group(1))  # 각주 번호가 붙은 '20043' 같은 값 처리
        if current_year is None:
            continue

        fcd_val = ws.cell_value(r, fcd_col)
        if not isinstance(fcd_val, float):
            continue

        month_str = str(month_cell).strip().lower()[:3]
        if month_str in _MONTHS:
            month = _MONTHS[month_str]
        elif year_cell != "":
            month = 12
        else:
            continue

        period = f"{current_year}-{month:02d}"
        rows.append({
            "country_code": country_code, "year": current_year, "period": period,
            "indicator": INDICATOR, "value": round(float(fcd_val), 4), "updated_at": now,
        })

        td_val = ws.cell_value(r, td_col) if td_col else None
        if isinstance(td_val, float):
            rows.append({
                "country_code": country_code, "year": current_year, "period": period,
                "indicator": "TD", "value": round(float(td_val), 4), "updated_at": now,
            })

    return pd.DataFrame(rows)


def render(target: dict) -> pd.DataFrame:
    country_code = target["country_code"]

    links = _collect_efs_links()
    logger.info("[%s] EFS xlsx %d개 발견", country_code, len(links))

    file_frames = []

    try:
        response = requests.get(LEGACY_XLS_URL, headers=_HEADERS, timeout=60)
        response.raise_for_status()
        legacy_df = _parse_legacy_xls(response.content, country_code)
        logger.info("[%s] 레거시 xls(2004~2014) %d행", country_code, len(legacy_df))
        if not legacy_df.empty:
            file_frames.append((legacy_df["period"].max(), legacy_df))
    except Exception:
        logger.warning("[%s] 레거시 xls 다운로드/파싱 실패, 스킵: %s", country_code, LEGACY_XLS_URL)

    for url in links:
        try:
            response = requests.get(url, headers=_HEADERS, timeout=60)
            response.raise_for_status()
        except Exception:
            logger.warning("[%s] 다운로드 실패, 스킵: %s", country_code, url)
            continue

        df = _parse_workbook(response.content, country_code)
        if df.empty:
            continue
        file_frames.append((df["period"].max(), df))

    if not file_frames:
        return pd.DataFrame(columns=["country_code", "year", "period", "indicator", "value", "updated_at"])

    # 각 파일이 커버하는 마지막 기간(발행월) 오름차순으로 합쳐, 겹치는 기간은
    # 더 최근에 발행된 파일 쪽 값이 남도록 한다(keep='last').
    file_frames.sort(key=lambda x: x[0])
    merged = pd.concat([df for _, df in file_frames], ignore_index=True)
    merged = merged.drop_duplicates(subset=["period", "indicator"], keep="last")
    return merged.sort_values(["indicator", "period"]).reset_index(drop=True)
