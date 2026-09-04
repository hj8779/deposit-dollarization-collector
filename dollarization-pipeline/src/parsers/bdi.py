"""Burundi: sheet 'Mensuelle', col A=date, col D='Dépôts en devises des résidents'.

TD (total deposits) = col B(Dépôts à vue, demand deposits) + col C(Dépôts à terme et d'épargne,
term/savings deposits) + col D(Dépôts en devises des résidents, FCD). This covers all resident
deposits regardless of currency; col E~K (Etablissements de Microfinances / interbank deposits,
central bank/government liabilities, etc.) are excluded."""

from datetime import datetime, timezone
from io import BytesIO

import openpyxl
import pandas as pd

from src.collectors.base import INDICATOR, INDICATOR_TD

FILE_URL = None  # uses target['source_url'] as-is


def parse(content: bytes, country_code: str) -> pd.DataFrame:
    wb = openpyxl.load_workbook(BytesIO(content), data_only=True)
    ws = wb["Mensuelle"]
    now = datetime.now(timezone.utc).isoformat()
    rows = []
    for r in range(8, ws.max_row + 1):
        date = ws.cell(row=r, column=1).value
        demand = ws.cell(row=r, column=2).value
        term_savings = ws.cell(row=r, column=3).value
        fcd = ws.cell(row=r, column=4).value
        if not isinstance(date, datetime) or fcd is None:
            continue
        period = f"{date.year}-{date.month:02d}"
        rows.append({
            "country_code": country_code,
            "year": date.year,
            "period": period,
            "indicator": INDICATOR,
            "value": float(fcd),
            "updated_at": now,
        })
        if demand is not None and term_savings is not None:
            rows.append({
                "country_code": country_code,
                "year": date.year,
                "period": period,
                "indicator": INDICATOR_TD,
                "value": round(float(demand) + float(term_savings) + float(fcd), 2),
                "updated_at": now,
            })
    return pd.DataFrame(rows)
