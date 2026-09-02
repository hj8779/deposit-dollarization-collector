"""Iraq: CBI "Key Financial Indicators" xlsx (linked from https://cbi.iq/news/view/94,
under-the-hood filename changes every publication so we scrape the current link).

Sheet1, "5- Money Supply" block:
- row "M2"                              -> Broad money
- row "Of which: Deposits component of M2" -> M2 minus currency outside banks;
  verified against Table 17 in the Annual Statistical Bulletin (matches exactly) ->
  used as TD (총예금, local+foreign currency deposits together).
- row "Foreign currency deposits" (subset of the above) -> FCD.

Column headers: row1 carries the year (forward-filled across that year's 12 month
columns, blank otherwise), row2 the month abbreviation. The sheet also has a few
weekly columns at the very end (row2 holds a date, not a month name) which we skip
since they carry no data on these rows.

Series starts 2003-12 (TD) / 2004-12 (FCD) — nothing before that in this source.
1991-2003 (the "Combined Historical Document" PDF covering that gap) is a scanned
multi-decade table that isn't reliably machine-parseable, so that range is left to
manual entry (see web/src/lib/manualUpdateCountries.ts: IRQ half_manual)."""

import re
from datetime import datetime, timezone
from io import BytesIO

import openpyxl
import pandas as pd
import requests
import urllib3

from src.collectors.base import INDICATOR, INDICATOR_TD

FILE_URL = "__RENDER__"

_NEWS_URL = "https://cbi.iq/news/view/94"
_XLSX_RE = re.compile(r"https?://cbi\.iq/static/uploads/up/file-[^\"'\s>]+\.xlsx")

_MONTHS = {
    "jan": 1, "feb": 2, "mar": 3, "march": 3, "apr": 4, "april": 4, "may": 5,
    "jun": 6, "june": 6, "jul": 7, "july": 7, "aug": 8, "august": 8, "sep": 9,
    "sept": 9, "september": 9, "oct": 10, "nov": 11, "november": 11,
    "dec": 12, "december": 12,
}

_M2_LABEL = "m2"
_DEPOSITS_LABEL = "deposits component of m2"
_FCD_LABEL = "foreign currency deposits"


def _month_num(value) -> int | None:
    if not isinstance(value, str):
        return None
    key = value.strip().lower().rstrip(".")
    return _MONTHS.get(key)


def _find_row(ws, label: str) -> int | None:
    for r in range(1, ws.max_row + 1):
        cell = ws.cell(row=r, column=2).value
        if isinstance(cell, str) and label in cell.strip().lower():
            return r
    return None


def render(target: dict) -> pd.DataFrame:
    urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)
    headers = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"}

    page = requests.get(_NEWS_URL, headers=headers, verify=False, timeout=30)
    page.raise_for_status()
    match = _XLSX_RE.search(page.text)
    if not match:
        return pd.DataFrame(columns=["country_code", "year", "period", "indicator", "value", "updated_at"])
    xlsx_url = match.group(0)

    resp = requests.get(xlsx_url, headers=headers, verify=False, timeout=60)
    resp.raise_for_status()
    return parse(resp.content, target["country_code"])


def parse(content: bytes, country_code: str) -> pd.DataFrame:
    wb = openpyxl.load_workbook(BytesIO(content), data_only=True)
    ws = wb["Sheet1"]

    m2_row = _find_row(ws, _M2_LABEL)
    deposits_row = _find_row(ws, _DEPOSITS_LABEL)
    fcd_row = _find_row(ws, _FCD_LABEL)
    if deposits_row is None and fcd_row is None:
        return pd.DataFrame(columns=["country_code", "year", "period", "indicator", "value", "updated_at"])

    year = None
    now = datetime.now(timezone.utc).isoformat()
    rows = []
    for c in range(1, ws.max_column + 1):
        year_cell = ws.cell(row=1, column=c).value
        if isinstance(year_cell, int):
            year = year_cell
        month_cell = ws.cell(row=2, column=c).value
        month = _month_num(month_cell)
        if year is None or month is None:
            continue
        period = f"{year}-{month:02d}"

        td_value = ws.cell(row=deposits_row, column=c).value if deposits_row else None
        if td_value is not None:
            rows.append({
                "country_code": country_code,
                "year": year,
                "period": period,
                "indicator": INDICATOR_TD,
                "value": float(td_value),
                "updated_at": now,
            })

        fcd_value = ws.cell(row=fcd_row, column=c).value if fcd_row else None
        if fcd_value is not None:
            rows.append({
                "country_code": country_code,
                "year": year,
                "period": period,
                "indicator": INDICATOR,
                "value": float(fcd_value),
                "updated_at": now,
            })

    return pd.DataFrame(rows)
