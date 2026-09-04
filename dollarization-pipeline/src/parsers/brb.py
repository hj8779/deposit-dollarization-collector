"""Barbados: stitches together two sources.

1) 1989-01~2011-12: from the 'commercial-banks-deposit-liabilities-2' page's
   'HISTORY 4 Commercial Banks DEPOSITS M 1989-2011.xlsx', sheet 'HISTORY 14'
   ('TOTAL DEPOSITS BY DEPOSITORS'); column J 'Deposits in Foreign Currency' is FCD, column K
   'Total Deposits' is TD (column K = column J + column I 'Total Domestic Deposits'). At first
   this file's leading columns (broken down by depositor: Government/Statutory Bodies/... and
   Demand·Time·Savings) were the only thing examined, leading to the mistaken conclusion that
   there was no breakdown by currency — but columns J/K already had the exact currency-level
   totals computed (pointed out by the user).

2) 2012-01~present: from the 'news/statistics-1' post listing, each month's 'Commercial Banks
   Assets and Liabilities {Month} {Year}' post contains an xlsx (cdn.centralbank.org.bb,
   filename containing 'BANKSASSETSLIABS'), 2nd tab 'BANKS - Liabilities'. Every file contains
   the entire cumulative time series from 2012-01 through that month, so only the latest post
   needs to be fetched (listing is sorted newest-first; the first matching link in the DOM is
   used). FCD = the sum of the 'Foreign Currency' column across the TRANSFERABLE DEPOSITS +
   OTHER DEPOSITS categories (regardless of whether Included/Excluded from Broad Money), with
   columns located dynamically via the row-5 (major category) / row-7 (currency) header text
   (per the no-fixed-index rule).

The boundary between the two sources (2011-12 -> 2012-01) lines up exactly with no overlapping
month, so they are simply concatenated.
"""

import re
from datetime import datetime, timezone

import pandas as pd
import requests

from src.collectors.base import INDICATOR
from src.utils.logger import get_logger

logger = get_logger(__name__)

FILE_URL = "__RENDER__"

HISTORY_LIST_URL = "https://www.centralbank.org.bb/news/commercial-banks/commercial-banks-deposit-liabilities-2"
LIST_URL = "https://www.centralbank.org.bb/news/statistics-1"
_HEADERS = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0 Safari/537.36"}

_POST_LINK_RE = re.compile(r"/statistics-1/commercial-banks-assets-and-liabilities-[a-z0-9-]+$", re.I)
_HISTORY_XLSX_RE = re.compile(r'https://cdn\.centralbank\.org\.bb/documents/[^"\']*HISTORY-4-Commercial-Banks-DEPOSITS[^"\']*\.xlsx', re.I)

_HISTORY_SHEET = "HISTORY 14"
_HISTORY_DATE_COL = 1
_HISTORY_FC_HEADER_RE = re.compile(r"deposits in foreign currency", re.I)
_HISTORY_TOTAL_HEADER_RE = re.compile(r"^total deposits\s*$", re.I)
_HISTORY_HEADER_ROW = 7

_SHEET_NAME = "BANKS - Liabilities"
_GROUP_ROW = 5
_CURRENCY_ROW = 7
_DEPOSIT_GROUP_RE = re.compile(r"deposit", re.I)
_DATE_COL = 1


def parse(content: bytes, country_code: str) -> pd.DataFrame:
    raise NotImplementedError("BRB is handled via render() (needs to fetch and stitch together two sources)")


def _find_latest_post_url() -> str | None:
    from playwright.sync_api import sync_playwright

    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page()
        page.goto(LIST_URL, timeout=60000)
        page.wait_for_timeout(2000)
        hrefs = page.eval_on_selector_all("a[href]", "els => els.map(e => e.href)")
        browser.close()

    for href in hrefs:
        if _POST_LINK_RE.search(href):
            return href
    return None


def _find_link(page_url: str, pattern: re.Pattern) -> str | None:
    from playwright.sync_api import sync_playwright

    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page()
        page.goto(page_url, timeout=60000)
        page.wait_for_timeout(2000)
        hrefs = page.eval_on_selector_all("a[href]", "els => els.map(e => e.href)")
        browser.close()

    for href in hrefs:
        if pattern.search(href):
            return href
    return None


def _find_xlsx_url(post_url: str) -> str | None:
    response = requests.get(post_url, headers=_HEADERS, timeout=30)
    response.raise_for_status()
    m = re.search(r'https://cdn\.centralbank\.org\.bb/documents/[^"\']+\.xlsx', response.text)
    return m.group(0) if m else None


def _parse_history(content: bytes, country_code: str) -> pd.DataFrame:
    import openpyxl
    from io import BytesIO

    now = datetime.now(timezone.utc).isoformat()
    wb = openpyxl.load_workbook(BytesIO(content), data_only=True)
    ws = wb[_HISTORY_SHEET]

    fc_col = td_col = None
    for c in range(1, ws.max_column + 1):
        header = ws.cell(row=_HISTORY_HEADER_ROW, column=c).value
        if not isinstance(header, str):
            continue
        if _HISTORY_FC_HEADER_RE.search(header):
            fc_col = c
        elif _HISTORY_TOTAL_HEADER_RE.match(header.strip()):
            td_col = c

    if fc_col is None:
        logger.warning("[%s] Could not find the 'Deposits in Foreign Currency' column in the history file", country_code)
        return pd.DataFrame(columns=["country_code", "year", "period", "indicator", "value", "updated_at"])

    rows = []
    for r in range(_HISTORY_HEADER_ROW + 1, ws.max_row + 1):
        date_val = ws.cell(row=r, column=_HISTORY_DATE_COL).value
        if not isinstance(date_val, datetime):
            continue
        fc_val = ws.cell(row=r, column=fc_col).value
        if not isinstance(fc_val, (int, float)):
            continue

        period = f"{date_val.year}-{date_val.month:02d}"
        rows.append({
            "country_code": country_code, "year": date_val.year, "period": period,
            "indicator": INDICATOR, "value": round(float(fc_val), 2), "updated_at": now,
        })

        td_val = ws.cell(row=r, column=td_col).value if td_col else None
        if isinstance(td_val, (int, float)):
            rows.append({
                "country_code": country_code, "year": date_val.year, "period": period,
                "indicator": "TD", "value": round(float(td_val), 2), "updated_at": now,
            })

    return pd.DataFrame(rows)


def _fc_deposit_columns(ws) -> list[int]:
    # Find the columns where a major category starts in row 5, then within each such span collect only the columns where row 7 is 'Foreign Currency'.
    group_starts = [
        c for c in range(2, ws.max_column + 1)
        if isinstance(ws.cell(row=_GROUP_ROW, column=c).value, str) and ws.cell(row=_GROUP_ROW, column=c).value.strip()
    ]

    fc_cols = []
    for i, start in enumerate(group_starts):
        label = ws.cell(row=_GROUP_ROW, column=start).value.strip()
        end = group_starts[i + 1] - 1 if i + 1 < len(group_starts) else ws.max_column
        if not _DEPOSIT_GROUP_RE.search(label):
            continue
        for c in range(start, end + 1):
            currency = ws.cell(row=_CURRENCY_ROW, column=c).value
            if isinstance(currency, str) and currency.strip().lower() == "foreign currency":
                fc_cols.append(c)
    return fc_cols


def _deposit_columns(ws) -> list[int]:
    group_starts = [
        c for c in range(2, ws.max_column + 1)
        if isinstance(ws.cell(row=_GROUP_ROW, column=c).value, str) and ws.cell(row=_GROUP_ROW, column=c).value.strip()
    ]
    cols = []
    for i, start in enumerate(group_starts):
        label = ws.cell(row=_GROUP_ROW, column=start).value.strip()
        end = group_starts[i + 1] - 1 if i + 1 < len(group_starts) else ws.max_column
        if _DEPOSIT_GROUP_RE.search(label):
            cols.extend(range(start, end + 1))
    return cols


def _parse_recent(content: bytes, country_code: str) -> pd.DataFrame:
    import openpyxl
    from io import BytesIO

    now = datetime.now(timezone.utc).isoformat()
    wb = openpyxl.load_workbook(BytesIO(content), data_only=True)
    ws = wb[_SHEET_NAME]

    fc_cols = _fc_deposit_columns(ws)
    all_cols = _deposit_columns(ws)
    if not fc_cols:
        logger.warning("[%s] Could not find the 'Foreign Currency' deposit column", country_code)
        return pd.DataFrame(columns=["country_code", "year", "period", "indicator", "value", "updated_at"])

    rows = []
    for r in range(_CURRENCY_ROW + 1, ws.max_row + 1):
        date_val = ws.cell(row=r, column=_DATE_COL).value
        if not isinstance(date_val, datetime):
            continue

        fc_values = [ws.cell(row=r, column=c).value for c in fc_cols]
        all_values = [ws.cell(row=r, column=c).value for c in all_cols]
        if not all(isinstance(v, (int, float)) for v in fc_values):
            continue

        period = f"{date_val.year}-{date_val.month:02d}"
        rows.append({
            "country_code": country_code, "year": date_val.year, "period": period,
            "indicator": INDICATOR, "value": round(sum(fc_values), 2), "updated_at": now,
        })
        if all(isinstance(v, (int, float)) for v in all_values):
            rows.append({
                "country_code": country_code, "year": date_val.year, "period": period,
                "indicator": "TD", "value": round(sum(all_values), 2), "updated_at": now,
            })

    return pd.DataFrame(rows)


def render(target: dict) -> pd.DataFrame:
    country_code = target["country_code"]
    frames = []

    history_url = _find_link(HISTORY_LIST_URL, _HISTORY_XLSX_RE)
    if history_url:
        response = requests.get(history_url, headers=_HEADERS, timeout=60)
        response.raise_for_status()
        hist_df = _parse_history(response.content, country_code)
        logger.info("[%s] History (1989-2011): %d rows", country_code, len(hist_df))
        if not hist_df.empty:
            frames.append(hist_df)
    else:
        logger.warning("[%s] Could not find the history xlsx link", country_code)

    post_url = _find_latest_post_url()
    if post_url:
        xlsx_url = _find_xlsx_url(post_url)
        if xlsx_url:
            logger.info("[%s] %s -> %s", country_code, post_url, xlsx_url)
            response = requests.get(xlsx_url, headers=_HEADERS, timeout=60)
            response.raise_for_status()
            recent_df = _parse_recent(response.content, country_code)
            logger.info("[%s] Recent (2012~): %d rows", country_code, len(recent_df))
            if not recent_df.empty:
                frames.append(recent_df)
        else:
            logger.warning("[%s] Could not find an xlsx link inside the post: %s", country_code, post_url)
    else:
        logger.warning("[%s] Could not find the latest 'Commercial Banks Assets and Liabilities' post", country_code)

    if not frames:
        return pd.DataFrame(columns=["country_code", "year", "period", "indicator", "value", "updated_at"])

    df = pd.concat(frames, ignore_index=True)
    df = df.drop_duplicates(subset=["period", "indicator"], keep="last")
    return df.sort_values(["indicator", "period"]).reset_index(drop=True)
