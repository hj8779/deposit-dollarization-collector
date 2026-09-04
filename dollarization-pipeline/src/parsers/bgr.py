"""Bulgaria: BNB (Bulgarian National Bank) 'Short Monetary Survey' XLSX, sheet 'MS_short'
(first tab). A horizontally expanding time series (row2=date, col2~=monthly). The
currency breakdown (BGN/foreign currency) appears as a pair of sub-rows, 'in BGN' /
'in foreign currency', repeated under multiple categories (the same labels show up under
both asset items and liability/deposit items).

For FCD/TD, only the categories that correspond to 'deposits' must be selected (asset-side
categories like FOREIGN ASSETS, DOMESTIC CREDIT, CLAIMS ON..., etc. also have 'in foreign
currency' sub-rows, and summing everything would be wrong). Rule: only aggregate when the
immediately preceding parent (non-indented) category label contains 'deposit'. Exactly 4
categories match this condition:
    Overnight deposits
    Deposits with agreed maturity up to 2 years
    Deposits redeemable at notice up to 3 months
    Deposits with agreed maturity over 2 years and deposits redeemable at notice over 3 months
FCD = sum of 'in foreign currency' values across these 4 categories
TD  = sum of 'in BGN' + 'in foreign currency' values across these 4 categories (equals the sum
      of category totals)
"""

from datetime import datetime, timezone
from io import BytesIO

import openpyxl
import pandas as pd

from src.collectors.base import INDICATOR

FILE_URL = "https://www.bnb.bg/bnbweb/groups/public/documents/bnb_download/s_ms_monetarystatistics_all_en.xlsx"

_SHEET_NAME = "MS_short"
_HEADER_ROW = 2
_FIRST_DATA_COL = 2


def parse(content: bytes, country_code: str) -> pd.DataFrame:
    now = datetime.now(timezone.utc).isoformat()
    wb = openpyxl.load_workbook(BytesIO(content), data_only=True)
    ws = wb[_SHEET_NAME]

    # 1) Find the row numbers of 'in BGN'/'in foreign currency' rows under 'deposit' categories.
    bgn_rows, fc_rows = [], []
    current_parent = ""
    for r in range(3, ws.max_row + 1):
        label = ws.cell(row=r, column=1).value
        if not isinstance(label, str) or not label.strip():
            continue
        stripped = label.strip().lower()
        if stripped == "in bgn":
            if "deposit" in current_parent:
                bgn_rows.append(r)
        elif stripped == "in foreign currency":
            if "deposit" in current_parent:
                fc_rows.append(r)
        else:
            current_parent = stripped

    if not fc_rows:
        return pd.DataFrame(columns=["country_code", "year", "period", "indicator", "value", "updated_at"])

    # 2) Sum the values of the rows above, per date column.
    rows = []
    for c in range(_FIRST_DATA_COL, ws.max_column + 1):
        date_val = ws.cell(row=_HEADER_ROW, column=c).value
        if not isinstance(date_val, datetime):
            continue

        fc_values = [ws.cell(row=r, column=c).value for r in fc_rows]
        bgn_values = [ws.cell(row=r, column=c).value for r in bgn_rows]
        if not all(isinstance(v, (int, float)) for v in fc_values + bgn_values):
            continue

        fcd = sum(fc_values)
        td = fcd + sum(bgn_values)
        period = f"{date_val.year}-{date_val.month:02d}"

        rows.append({
            "country_code": country_code, "year": date_val.year, "period": period,
            "indicator": INDICATOR, "value": round(fcd, 2), "updated_at": now,
        })
        rows.append({
            "country_code": country_code, "year": date_val.year, "period": period,
            "indicator": "TD", "value": round(td, 2), "updated_at": now,
        })

    return pd.DataFrame(rows)
