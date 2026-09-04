"""New Zealand: RBNZ "Registered banks' balance sheet" (HS10) xlsx.

https://www.rbnz.govt.nz/statistics/series/registered-banks/banks-balance-sheet
https://www.rbnz.govt.nz/-/media/project/sites/rbnz/files/statistics/series/l-s/s10/hs10m.xlsx

rbnz.govt.nz is protected by a Cloudflare JS challenge ("Just a moment...") — a
direct requests fetch only gets a 403 (challenge HTML). Doing a plain goto() to the
URL with Playwright lets the browser pass the challenge, but a file download then
starts immediately, so goto() itself throws a "Download is starting" exception;
this is caught by wrapping the call in page.expect_download()
(confirmed that this single goto is sufficient to pass — no separate page visit or
cookie transplant needed).

'Data' sheet, header rows 1-5 (row 2 = series description), monthly data from row 6,
column A = date. Column F 'C1. Deposits (NZD)' = local-currency deposits, column G
'C2. Deposits (FX)' = foreign-currency deposits = FCD.
TD = C1 + C2. Monthly, from 2016-12. Unit: NZD million."""

from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO

import openpyxl
import pandas as pd

from src.collectors.base import INDICATOR, INDICATOR_TD
from src.utils.logger import get_logger

logger = get_logger(__name__)

FILE_URL = "__RENDER__"
_XLSX_URL = (
    "https://www.rbnz.govt.nz/-/media/project/sites/rbnz/files/statistics/series/l-s/s10/hs10m.xlsx"
)
_USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/122.0.0.0 Safari/537.36"
)

_NZD_COL = 6  # F: C1. Deposits (NZD)
_FX_COL = 7   # G: C2. Deposits (FX)


def parse(content: bytes, country_code: str) -> pd.DataFrame:
    raise NotImplementedError("NZL bypasses Cloudflare via render() to fetch the xlsx")


def _empty() -> pd.DataFrame:
    return pd.DataFrame(columns=["country_code", "year", "period", "indicator", "value", "updated_at"])


def _download_via_browser() -> bytes | None:
    from playwright.sync_api import sync_playwright

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True, args=["--disable-blink-features=AutomationControlled"])
        try:
            context = browser.new_context(user_agent=_USER_AGENT, accept_downloads=True)
            page = context.new_page()
            try:
                with page.expect_download(timeout=30000) as dl_info:
                    try:
                        page.goto(_XLSX_URL, timeout=30000)
                    except Exception:
                        pass  # goto throwing "Download is starting" is the expected path
                download = dl_info.value
                path = download.path()
                if not path:
                    return None
                with open(path, "rb") as f:
                    return f.read()
            except Exception:
                logger.warning("[NZL] Playwright download failed", exc_info=True)
                return None
        finally:
            browser.close()


def render(target: dict) -> pd.DataFrame:
    country_code = target["country_code"]
    content = _download_via_browser()
    if not content:
        return _empty()

    wb = openpyxl.load_workbook(BytesIO(content), data_only=True)
    ws = wb["Data"]

    now = datetime.now(timezone.utc).isoformat()
    rows = []
    for r in range(6, ws.max_row + 1):
        date = ws.cell(row=r, column=1).value
        if not isinstance(date, datetime):
            continue
        nzd = ws.cell(row=r, column=_NZD_COL).value
        fx = ws.cell(row=r, column=_FX_COL).value
        if not isinstance(nzd, (int, float)) or not isinstance(fx, (int, float)):
            continue

        fcd = float(fx)
        td = float(nzd) + fcd
        if td <= 0:
            continue

        period = f"{date.year}-{date.month:02d}"
        ratio = round((fcd / td) * 100, 4)
        for indicator, value in ((INDICATOR, round(fcd, 4)), (INDICATOR_TD, round(td, 4)), ("FCD_TD_RATIO", ratio)):
            rows.append({
                "country_code": country_code, "year": date.year, "period": period,
                "indicator": indicator, "value": value, "updated_at": now,
            })

    if not rows:
        return _empty()

    out = (
        pd.DataFrame(rows)
        .drop_duplicates(subset=["period", "indicator"], keep="last")
        .sort_values(["period", "indicator"])
        .reset_index(drop=True)
    )
    logger.info("[%s] %d rows (%s~%s)", country_code, len(out), out["period"].min(), out["period"].max())
    return out
