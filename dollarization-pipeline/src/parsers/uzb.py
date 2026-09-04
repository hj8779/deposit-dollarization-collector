"""Uzbekistan: CBU "Monetary aggregates" xlsx.

The xlsx link is scraped from https://cbu.uz/en/statistics/dks/4320467/ (Monetary and
financial statistics > Monetary aggregates), since the document path can change.

The English / Russian ("Денежные агрегаты") / Uzbek ("Pul agregatlari") pages each use a
different content ID, but the data is identical across all three and starts at 2013-02
(verified by downloading and cross-checking all three language versions directly) - it
appears the CBU statistic itself has no time series before 2013 (presumably because the
currency-based M2 breakdown, per the 2016 IMF Monetary and Financial Statistics Manual, was
only applied retroactively starting from 2013).

Sheet 'Broad money': row 7 is the column-index header (formulas like 2=3+8 indicate column
composition), monthly data starts at row 8. Column B (2) = Broad money M2 total = TD.
Column H (8) = 'Foreign currency deposits in national currency equivalent' = FCD. Unit:
billion UZS."""

from __future__ import annotations

import re
from datetime import datetime, timezone
from io import BytesIO

import openpyxl
import pandas as pd
import requests

from src.collectors.base import INDICATOR, INDICATOR_TD
from src.utils.logger import get_logger

logger = get_logger(__name__)

FILE_URL = "__RENDER__"

_LIST_PAGE = "https://cbu.uz/en/statistics/dks/4320467/"
_XLSX_LINK_RE = re.compile(r'href="([^"]*\.xlsx)"', re.I)
_HEADERS = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/122.0.0.0 Safari/537.36"}

_TD_COL = 2  # Broad money (M2)
_FCD_COL = 8  # Foreign currency deposits in national currency equivalent
_FIRST_DATA_ROW = 8


def parse(content: bytes, country_code: str) -> pd.DataFrame:
    raise NotImplementedError("UZB fetches the xlsx via render()")


def _empty() -> pd.DataFrame:
    return pd.DataFrame(columns=["country_code", "year", "period", "indicator", "value", "updated_at"])


def _find_xlsx_url() -> str | None:
    resp = requests.get(_LIST_PAGE, headers=_HEADERS, timeout=30)
    resp.raise_for_status()
    match = _XLSX_LINK_RE.search(resp.text)
    if not match:
        return None
    href = match.group(1)
    return href if href.startswith("http") else "https://cbu.uz" + href


def render(target: dict) -> pd.DataFrame:
    country_code = target["country_code"]
    xlsx_url = _find_xlsx_url()
    if not xlsx_url:
        logger.warning("[%s] xlsx link not found on listing page", country_code)
        return _empty()

    resp = requests.get(xlsx_url, headers=_HEADERS, timeout=60)
    resp.raise_for_status()
    wb = openpyxl.load_workbook(BytesIO(resp.content), data_only=True)
    ws = wb[wb.sheetnames[0]]

    now = datetime.now(timezone.utc).isoformat()
    rows = []
    for r in range(_FIRST_DATA_ROW, ws.max_row + 1):
        date = ws.cell(row=r, column=1).value
        if not isinstance(date, datetime):
            continue
        td = ws.cell(row=r, column=_TD_COL).value
        fcd = ws.cell(row=r, column=_FCD_COL).value
        if not isinstance(td, (int, float)) or not isinstance(fcd, (int, float)) or td <= 0:
            continue

        period = f"{date.year}-{date.month:02d}"
        ratio = round((float(fcd) / float(td)) * 100, 4)
        for indicator, value in ((INDICATOR, round(float(fcd), 4)), (INDICATOR_TD, round(float(td), 4)), ("FCD_TD_RATIO", ratio)):
            rows.append({
                "country_code": country_code, "year": date.year, "period": period,
                "indicator": indicator, "value": value, "updated_at": now,
            })

    if not rows:
        return _empty()

    out = (
        pd.DataFrame(rows)
        .drop_duplicates(subset=["period", "indicator"], keep="last")
        .sort_values(["period", "indicator"])
        .reset_index(drop=True)
    )
    logger.info("[%s] %d rows (%s~%s)", country_code, len(out), out["period"].min(), out["period"].max())
    return out
