"""Source 탭과 Target 탭을 alpha3_code 기준으로 병합해 JSON으로 출력한다.

사용법:
    python3 scripts/merge_source_target.py [입력.xlsx] [출력.json]

기본값:
    입력 = Data_dolrat_raw.xlsx
    출력 = schema/data.json
"""

import json
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent


def load_source(ws):
    rows = {}
    for r in range(2, ws.max_row + 1):
        name, alpha2, alpha3, numeric, url1, url2, notes = (
            ws.cell(row=r, column=c).value for c in range(1, 8)
        )
        if not alpha3:
            continue
        rows[alpha3] = {
            "country_name": name,
            "alpha2_code": alpha2,
            "alpha3_code": alpha3,
            "numeric_code": int(numeric) if numeric is not None else None,
            "source": {
                "source_url": url1,
                "source_url_2": url2,
                "notes": notes,
            },
            "target": None,
        }
    return rows


def load_target(ws):
    rows = {}
    for r in range(2, ws.max_row + 1):
        alpha3, data_type, url, freq, requires_js, notes = (
            ws.cell(row=r, column=c).value for c in range(1, 7)
        )
        if not alpha3:
            continue
        rows[alpha3] = {
            "data_type": data_type,
            "direct_download_url": url,
            "update_frequency": freq,
            "requires_js": bool(requires_js) if requires_js is not None else None,
            "notes": notes,
        }
    return rows


def main():
    src_path = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "Data_dolrat_raw.xlsx"
    out_path = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / "schema" / "data.json"

    wb = openpyxl.load_workbook(src_path, data_only=True)
    merged = load_source(wb["Source"])
    target = load_target(wb["Target"])

    for alpha3, target_row in target.items():
        if alpha3 in merged:
            merged[alpha3]["target"] = target_row

    result = list(merged.values())

    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"wrote {len(result)} records to {out_path}")


if __name__ == "__main__":
    main()
